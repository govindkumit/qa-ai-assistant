from pathlib import Path

from app.pdf_loader import extract_text_from_pdf


KNOWLEDGE_DIR = Path("knowledge")


def load_txt(file_path):
    text = file_path.read_text(
        encoding="utf-8"
    )

    return [
        {
            "source": file_path.name,
            "text": text
        }
    ]


def load_pdf(file_path):
    pages = extract_text_from_pdf(file_path)

    documents = []

    for page in pages:

        documents.append(
            {
                "source": file_path.name,
                "text": page["text"],
                "page_number": page["page_number"]
            }
        )

    return documents


def load_docx(file_path):

    from docx import Document

    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    text = "\n".join(paragraphs)

    if not text:
        return []

    return [
        {
            "source": file_path.name,
            "text": text
        }
    ]


def load_xlsx(file_path):

    from openpyxl import load_workbook

    workbook = load_workbook(
        file_path,
        data_only=True
    )

    documents = []

    for sheet in workbook.worksheets:

        rows = []

        for row in sheet.iter_rows(
            values_only=True
        ):

            values = []

            for value in row:

                if value is not None:
                    values.append(str(value))

            if values:
                rows.append(" | ".join(values))

        if rows:

            documents.append(
                {
                    "source": file_path.name,
                    "text": "\n".join(rows),
                    "sheet_name": sheet.title
                }
            )

    return documents


def load_documents():

    documents = []

    KNOWLEDGE_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    supported_extensions = {
        ".txt",
        ".pdf",
        ".docx",
        ".xlsx"
    }

    for file_path in KNOWLEDGE_DIR.iterdir():

        if not file_path.is_file():
            continue

        extension = file_path.suffix.lower()

        if extension not in supported_extensions:
            continue

        try:

            if extension == ".txt":
                loaded = load_txt(file_path)

            elif extension == ".pdf":
                loaded = load_pdf(file_path)

            elif extension == ".docx":
                loaded = load_docx(file_path)

            elif extension == ".xlsx":
                loaded = load_xlsx(file_path)

            documents.extend(loaded)

            print(
                f"Loaded: {file_path.name}"
            )

        except Exception as error:

            print(
                f"Failed to load "
                f"{file_path.name}: {error}"
            )

    return documents


def chunk_document(
    text,
    chunk_size=500
):

    words = text.split()

    chunks = []

    for start in range(
        0,
        len(words),
        chunk_size
    ):

        chunk = " ".join(
            words[
                start:start + chunk_size
            ]
        )

        if chunk.strip():
            chunks.append(chunk)

    return chunks


if __name__ == "__main__":

    documents = load_documents()

    total_chunks = 0

    for document in documents:

        chunks = chunk_document(
            document["text"]
        )

        print(
            f"\nSource: "
            f"{document['source']}"
        )

        print(
            f"Chunks: {len(chunks)}"
        )

        total_chunks += len(chunks)

    print(
        f"\nTotal documents/sections: "
        f"{len(documents)}"
    )

    print(
        f"Total chunks: {total_chunks}"
    )